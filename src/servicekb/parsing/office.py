"""PDF / Word / Excel 解析。

PDF 是工业文档里最难啃的格式，常见三种情况：
1. **原生文本 PDF**（Word 导出的）：pypdf 能直接抽文本，但丢失版面结构，需要启发式恢复标题。
2. **扫描件 PDF**：没有文本层，必须 OCR。pypdf 抽出来是空的 → 本模块会报错提示用 docling / MinerU。
3. **复杂版面**（多栏、跨页表格、中文字体编码损坏）：需要版面分析模型。

本项目默认只做第 1 种（零依赖、够快）；设置 ``SERVICEKB_PDF_PARSER=docling`` 并安装
``pip install -e ".[docling]"`` 后，PDF 会先经 docling 转成 Markdown 再解析。
MinerU 更适合作为独立的解析服务部署在 GPU 机器上（注意其许可证的营收条款，见 docs/07_repo_survey.md）。
"""

from __future__ import annotations

import io
import re

from servicekb.config import get_settings
from servicekb.models import Block, ParsedDocument, SourceRecord
from servicekb.parsing.markdown import parse_markdown_text
from servicekb.processing.text import normalize_text

# --------------------------------------------------------------------------- PDF

# 标题识别规则（借鉴 RAGFlow 的 BULLET_PATTERN）：编号 + 短句 + 不以句号结尾
_HEADING_PATTERNS: list[tuple[re.Pattern[str], int | None]] = [
    (re.compile(r"^第[一二三四五六七八九十百零\d]+[章]\s*\S"), 1),
    (re.compile(r"^第[一二三四五六七八九十百零\d]+[节]\s*\S"), 2),
    (re.compile(r"^Chapter\s+\d+\b", re.IGNORECASE), 1),
    (re.compile(r"^(\d+(?:\.\d+){0,3})\s+\S"), None),  # 1 / 1.2 / 1.2.3 → 层级 = 点数 + 1
    (re.compile(r"^[（(][一二三四五六七八九十]+[)）]\s*\S"), 3),
]
_SENT_END = re.compile(r"[。．.!！?？;；:：,，]$")


def _heading_level(line: str) -> int | None:
    if len(line) > 40 or _SENT_END.search(line):
        return None
    for pat, level in _HEADING_PATTERNS:
        m = pat.match(line)
        if m:
            if level is None:
                return min(m.group(1).count(".") + 1, 6)
            return level
    return None


def _join_lines(lines: list[str]) -> str:
    """PDF 抽出来的文本是按"视觉行"断开的，需要拼回段落：中文直接拼，英文加空格。"""
    out = ""
    for ln in lines:
        if not out:
            out = ln
        elif re.search(r"[一-鿿]$", out) or re.match(r"^[一-鿿]", ln):
            out += ln
        else:
            out += " " + ln
    return out


def _parse_pdf_with_docling(record: SourceRecord) -> ParsedDocument:
    from docling.document_converter import DocumentConverter  # 可选依赖，按需导入
    from docling_core.types.io import DocumentStream

    assert isinstance(record.content, bytes)
    stream = DocumentStream(name=record.external_id, stream=io.BytesIO(record.content))
    md = DocumentConverter().convert(stream).document.export_to_markdown()
    blocks, meta = parse_markdown_text(md)
    return ParsedDocument(record=record, blocks=blocks, metadata=meta)


def parse_pdf(record: SourceRecord) -> ParsedDocument:
    if get_settings().pdf_parser == "docling":
        return _parse_pdf_with_docling(record)

    from pypdf import PdfReader

    assert isinstance(record.content, bytes), "PDF content must be bytes"
    reader = PdfReader(io.BytesIO(record.content))
    blocks: list[Block] = []
    total_chars = 0
    for page_no, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        total_chars += len(text)
        para: list[str] = []
        # 末尾追加一个空行作为"哨兵"，保证最后一段也会被 flush
        for line in [*text.split("\n"), ""]:
            line = line.strip()
            level = _heading_level(line) if line else None
            if line and not level:
                para.append(line)
            # 遇到空行 / 标题 / 句末标点：结束当前段落
            if para and (not line or level or line[-1] in "。.!！?？"):
                blocks.append(Block(kind="paragraph", text=_join_lines(para), page=page_no))
                para = []
            if level:
                blocks.append(Block(kind="heading", text=line, level=level, page=page_no))

    if total_chars < 20 * max(len(reader.pages), 1):
        # 平均每页不到 20 个字符，几乎可以肯定是扫描件
        raise ValueError(
            f"{record.ref}: PDF has almost no text layer (scanned?). "
            "Use SERVICEKB_PDF_PARSER=docling or run MinerU/OCR first."
        )
    return ParsedDocument(record=record, blocks=blocks)


# --------------------------------------------------------------------------- Word


def parse_docx(record: SourceRecord) -> ParsedDocument:
    """按文档顺序遍历段落和表格。标题层级来自样式名 "Heading 1" / "标题 1"。"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    assert isinstance(record.content, bytes)
    d = docx.Document(io.BytesIO(record.content))
    blocks: list[Block] = []
    # python-docx 的 d.paragraphs / d.tables 是分开的，会打乱顺序；直接遍历 body 的子元素保证顺序
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = normalize_text(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"^(?:Heading|标题)\s*(\d)", style)
            if m:
                blocks.append(Block(kind="heading", text=text, level=int(m.group(1))))
            elif style.lower().startswith("title"):
                record.title = record.title or text
            elif "List" in style:
                blocks.append(Block(kind="list", text=text))
            else:
                blocks.append(Block(kind="paragraph", text=text))
        elif tag == "tbl":
            t = Table(child, d)
            rows = [[normalize_text(c.text) for c in r.cells] for r in t.rows]
            blocks.append(Block(kind="table", rows=rows))
    return ParsedDocument(record=record, blocks=blocks)


# --------------------------------------------------------------------------- Excel


def parse_xlsx(record: SourceRecord) -> ParsedDocument:
    """每个 sheet 一个标题 + 一张表。第一行非空行视为表头。"""
    from openpyxl import load_workbook

    assert isinstance(record.content, bytes)
    wb = load_workbook(io.BytesIO(record.content), read_only=True, data_only=True)
    blocks: list[Block] = []
    for ws in wb.worksheets:
        rows = [
            ["" if v is None else normalize_text(str(v)) for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        blocks.append(Block(kind="heading", text=ws.title, level=1))
        blocks.append(Block(kind="table", rows=rows))
    return ParsedDocument(record=record, blocks=blocks)
